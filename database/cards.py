import json
import subprocess
import os
import asyncio
from typing import List, Dict, Optional, Tuple
from .db import get_db
from config import CARD_TYPES

# ============================================
# АВТОМАТИЧЕСКИЙ БЭКАП
# ============================================

async def auto_backup():
    """Создать резервную копию базы данных (в фоновом режиме, без блокировки)"""
    asyncio.create_task(_run_backup_async())

async def _run_backup_async():
    """Запуск бэкапа в отдельном потоке (не блокирует event loop)"""
    try:
        # Используем run_in_executor для запуска в отдельном потоке
        loop = asyncio.get_event_loop()
        await loop.run_in_executor(None, _run_backup_sync)
    except Exception as e:
        print(f"⚠️ Ошибка при создании бэкапа: {e}")

def _run_backup_sync():
    """Синхронная функция для запуска в отдельном потоке"""
    try:
        import subprocess
        result = subprocess.run(
            ["python", "backup_db.py"],
            capture_output=True,
            text=True,
            cwd=os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        )
        if result.returncode == 0:
            print("✅ Автоматический бэкап создан")
        else:
            print(f"⚠️ Ошибка при создании бэкапа: {result.stderr}")
    except Exception as e:
        print(f"⚠️ Не удалось создать бэкап: {e}")

# ============================================
# УПРАВЛЕНИЕ КАРТОЧКАМИ (CRUD)
# ============================================

async def add_card(
    word: str,
    translation: str,
    language_from: str,
    language_to: str,
    card_type: str,
    created_by: int,
    image_url: str = None,
    wrong_hints: List[str] = None,
    topics: List[str] = None,
    alternative_translations: List[str] = None
) -> int:
    """Добавить новую карточку"""
    wrong_hints_json = json.dumps(wrong_hints) if wrong_hints else None
    topics_str = ','.join(topics) if topics else None
    alt_trans_str = ','.join(alternative_translations) if alternative_translations else None
    first_letter = word[0].upper() if word else None
    
    db = await get_db()
    cursor = await db.execute('''
        INSERT INTO cards (
            word, translation, language_from, language_to, 
            image_url, card_type, wrong_hints, topics, 
            alternative_translations, first_letter, created_by
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    ''', (word, translation, language_from, language_to, 
          image_url, card_type, wrong_hints_json, topics_str, 
          alt_trans_str, first_letter, created_by))
    await db.commit()
    await auto_backup()
    return cursor.lastrowid

async def get_card_by_id(card_id: int) -> Optional[Dict]:
    """Получить карточку по ID"""
    db = await get_db()
    async with db.execute('SELECT * FROM cards WHERE card_id = ?', (card_id,)) as cursor:
        row = await cursor.fetchone()
        if row:
            columns = ['card_id', 'word', 'translation', 'language_from', 
                      'language_to', 'image_url', 'card_type', 'wrong_hints', 
                      'topics', 'alternative_translations', 'first_letter', 
                      'created_by', 'created_at', 'updated_at']
            card_dict = dict(zip(columns, row))
            
            # Преобразуем JSON-поля обратно в списки
            if card_dict['wrong_hints']:
                card_dict['wrong_hints'] = json.loads(card_dict['wrong_hints'])
            else:
                card_dict['wrong_hints'] = []
                
            if card_dict['topics']:
                card_dict['topics'] = card_dict['topics'].split(',')
            else:
                card_dict['topics'] = []
                
            if card_dict['alternative_translations']:
                card_dict['alternative_translations'] = card_dict['alternative_translations'].split(',')
            else:
                card_dict['alternative_translations'] = []
                
            return card_dict
        return None

async def update_card(
    card_id: int,
    word: str = None,
    translation: str = None,
    language_from: str = None,
    language_to: str = None,
    card_type: str = None,
    image_url: str = None,
    wrong_hints: List[str] = None,
    topics: List[str] = None,
    alternative_translations: List[str] = None
) -> bool:
    """Обновить существующую карточку"""
    updates = []
    params = []
    
    if word is not None:
        updates.append("word = ?")
        params.append(word)
        updates.append("first_letter = ?")
        params.append(word[0].upper() if word else None)
    if translation is not None:
        updates.append("translation = ?")
        params.append(translation)
    if language_from is not None:
        updates.append("language_from = ?")
        params.append(language_from)
    if language_to is not None:
        updates.append("language_to = ?")
        params.append(language_to)
    if card_type is not None:
        updates.append("card_type = ?")
        params.append(card_type)
    if image_url is not None:
        updates.append("image_url = ?")
        params.append(image_url)
    if wrong_hints is not None:
        updates.append("wrong_hints = ?")
        params.append(json.dumps(wrong_hints))
    if topics is not None:
        updates.append("topics = ?")
        params.append(','.join(topics))
    if alternative_translations is not None:
        updates.append("alternative_translations = ?")
        params.append(','.join(alternative_translations))
    
    if not updates:
        return False
    
    updates.append("updated_at = CURRENT_TIMESTAMP")
    params.append(card_id)
    
    db = await get_db()
    await db.execute(f"UPDATE cards SET {', '.join(updates)} WHERE card_id = ?", params)
    await db.commit()
    await auto_backup()
    return True

async def delete_card(card_id: int, deleted_by: int) -> bool:
    """Удалить карточку и сохранить в лог"""
    card = await get_card_by_id(card_id)
    if not card:
        return False
    
    db = await get_db()
    # Сохраняем в лог удалённых
    await db.execute('''
        INSERT INTO deleted_cards_log (card_id, word, translation, deleted_by)
        VALUES (?, ?, ?, ?)
    ''', (card_id, card['word'], card['translation'], deleted_by))
    
    # Удаляем карточку
    await db.execute('DELETE FROM cards WHERE card_id = ?', (card_id,))
    
    # Удаляем прогресс пользователей по этой карточке
    await db.execute('DELETE FROM user_progress WHERE card_id = ?', (card_id,))
    
    await db.commit()
    await auto_backup()
    return True

async def get_all_cards(limit: int = 100, offset: int = 0) -> List[Dict]:
    """Получить все карточки с пагинацией"""
    db = await get_db()
    async with db.execute('''
        SELECT * FROM cards ORDER BY card_id DESC LIMIT ? OFFSET ?
    ''', (limit, offset)) as cursor:
        rows = await cursor.fetchall()
        columns = ['card_id', 'word', 'translation', 'language_from', 
                  'language_to', 'image_url', 'card_type', 'wrong_hints',
                  'topics', 'alternative_translations', 'first_letter', 
                  'created_by', 'created_at', 'updated_at']
        cards = []
        for row in rows:
            card_dict = dict(zip(columns, row))
            if card_dict['wrong_hints']:
                card_dict['wrong_hints'] = json.loads(card_dict['wrong_hints'])
            else:
                card_dict['wrong_hints'] = []
            if card_dict['topics']:
                card_dict['topics'] = card_dict['topics'].split(',')
            else:
                card_dict['topics'] = []
            if card_dict['alternative_translations']:
                card_dict['alternative_translations'] = card_dict['alternative_translations'].split(',')
            else:
                card_dict['alternative_translations'] = []
            cards.append(card_dict)
        return cards

async def search_cards(query_text: str) -> List[Dict]:
    """Поиск карточек по слову или переводу (админский)"""
    db = await get_db()
    async with db.execute('''
        SELECT * FROM cards 
        WHERE word LIKE ? OR translation LIKE ? 
        ORDER BY word ASC
    ''', (f'%{query_text}%', f'%{query_text}%')) as cursor:
        rows = await cursor.fetchall()
        columns = ['card_id', 'word', 'translation', 'language_from', 
                  'language_to', 'image_url', 'card_type', 'wrong_hints',
                  'topics', 'alternative_translations', 'first_letter', 
                  'created_by', 'created_at', 'updated_at']
        cards = []
        for row in rows:
            card_dict = dict(zip(columns, row))
            if card_dict['wrong_hints']:
                card_dict['wrong_hints'] = json.loads(card_dict['wrong_hints'])
            else:
                card_dict['wrong_hints'] = []
            if card_dict['topics']:
                card_dict['topics'] = card_dict['topics'].split(',')
            else:
                card_dict['topics'] = []
            if card_dict['alternative_translations']:
                card_dict['alternative_translations'] = card_dict['alternative_translations'].split(',')
            else:
                card_dict['alternative_translations'] = []
            cards.append(card_dict)
        return cards

# ============================================
# ФУНКЦИИ ДЛЯ СЛУЧАЙНЫХ КАРТОЧЕК
# ============================================

async def get_random_card(user_id: int, topic: str = None):
    """Получить случайную карточку для пользователя (без исключения)"""
    db = await get_db()
    if topic is None or topic == "Все слова":
        query = '''
            SELECT c.* FROM cards c
            LEFT JOIN user_progress up ON c.card_id = up.card_id AND up.user_id = ?
            WHERE (up.mastered IS NULL OR up.mastered = FALSE)
            ORDER BY RANDOM()
            LIMIT 1
        '''
        params = (user_id,)
    else:
        query = '''
            SELECT c.* FROM cards c
            LEFT JOIN user_progress up ON c.card_id = up.card_id AND up.user_id = ?
            WHERE (up.mastered IS NULL OR up.mastered = FALSE)
            AND c.topics LIKE ?
            ORDER BY RANDOM()
            LIMIT 1
        '''
        params = (user_id, f'%{topic}%')
    
    async with db.execute(query, params) as cursor:
        return await cursor.fetchone()

async def get_random_card_excluding(user_id: int, topic: str = None, exclude_card_id: int = None):
    """Получить случайную карточку, исключая указанную"""
    db = await get_db()
    
    if topic is None or topic == "Все слова":
        if exclude_card_id:
            query = '''
                SELECT c.* FROM cards c
                LEFT JOIN user_progress up ON c.card_id = up.card_id AND up.user_id = ?
                WHERE (up.mastered IS NULL OR up.mastered = FALSE)
                AND c.card_id != ?
                ORDER BY RANDOM()
                LIMIT 1
            '''
            params = (user_id, exclude_card_id)
        else:
            query = '''
                SELECT c.* FROM cards c
                LEFT JOIN user_progress up ON c.card_id = up.card_id AND up.user_id = ?
                WHERE (up.mastered IS NULL OR up.mastered = FALSE)
                ORDER BY RANDOM()
                LIMIT 1
            '''
            params = (user_id,)
    else:
        if exclude_card_id:
            query = '''
                SELECT c.* FROM cards c
                LEFT JOIN user_progress up ON c.card_id = up.card_id AND up.user_id = ?
                WHERE (up.mastered IS NULL OR up.mastered = FALSE)
                AND c.topics LIKE ?
                AND c.card_id != ?
                ORDER BY RANDOM()
                LIMIT 1
            '''
            params = (user_id, f'%{topic}%', exclude_card_id)
        else:
            query = '''
                SELECT c.* FROM cards c
                LEFT JOIN user_progress up ON c.card_id = up.card_id AND up.user_id = ?
                WHERE (up.mastered IS NULL OR up.mastered = FALSE)
                AND c.topics LIKE ?
                ORDER BY RANDOM()
                LIMIT 1
            '''
            params = (user_id, f'%{topic}%')
    
    async with db.execute(query, params) as cursor:
        return await cursor.fetchone()

# ============================================
# ФУНКЦИИ ДЛЯ ТЕМАТИК
# ============================================

async def get_all_topics() -> List[str]:
    """Получить все уникальные тематики из базы"""
    db = await get_db()
    async with db.execute('SELECT topics FROM cards WHERE topics IS NOT NULL AND topics != ""') as cursor:
        rows = await cursor.fetchall()
        topics = set()
        for row in rows:
            if row[0]:
                for t in row[0].split(','):
                    t = t.strip()
                    if t:
                        topics.add(t)
        return sorted(list(topics))

async def get_all_topics_paginated(page: int = 0, per_page: int = 10) -> Tuple[List[str], int]:
    """Получить тематики с пагинацией"""
    all_topics = await get_all_topics()
    total = len(all_topics)
    start = page * per_page
    end = start + per_page
    return all_topics[start:end], total

async def get_words_count_by_topic(topic: str) -> int:
    """Получить количество слов в тематике"""
    if topic == "Все слова":
        db = await get_db()
        async with db.execute('SELECT COUNT(*) FROM cards') as cursor:
            result = await cursor.fetchone()
            return result[0] if result else 0
    
    db = await get_db()
    async with db.execute('''
        SELECT COUNT(*) FROM cards 
        WHERE topics LIKE ?
    ''', (f'%{topic}%',)) as cursor:
        result = await cursor.fetchone()
        return result[0] if result else 0

async def get_all_topics_with_counts() -> List[Tuple[str, int]]:
    """Получить все тематики с количеством слов"""
    topics = await get_all_topics()
    result = []
    for topic in topics:
        count = await get_words_count_by_topic(topic)
        result.append((topic, count))
    return sorted(result, key=lambda x: x[0])

async def reset_user_progress_by_topic(user_id: int, topic: str) -> int:
    """Сбросить прогресс пользователя по словам из тематики"""
    if topic == "Все слова":
        from .progress import clear_user_stats
        await clear_user_stats(user_id)
        await auto_backup()
        return 0
    
    db = await get_db()
    # Получаем все карточки по тематике
    async with db.execute('''
        SELECT card_id FROM cards 
        WHERE topics LIKE ?
    ''', (f'%{topic}%',)) as cursor:
        rows = await cursor.fetchall()
        card_ids = [row[0] for row in rows]
    
    if not card_ids:
        return 0
    
    # Удаляем прогресс по этим карточкам
    placeholders = ','.join(['?'] * len(card_ids))
    query = f'''
        DELETE FROM user_progress 
        WHERE user_id = ? AND card_id IN ({placeholders})
    '''
    params = [user_id] + card_ids
    await db.execute(query, params)
    await db.commit()
    await auto_backup()
    return len(card_ids)

# ============================================
# ФУНКЦИИ ДЛЯ СЛОВАРЯ
# ============================================

async def get_words_by_letter(language_from: str, letter: str, page: int = 0, per_page: int = 5) -> Tuple[List[Dict], int]:
    """Получить слова по языку и первой букве с пагинацией"""
    db = await get_db()
    offset = page * per_page
    
    # Получаем общее количество
    query_count = '''
        SELECT COUNT(*) FROM cards 
        WHERE language_from = ? AND first_letter = ?
    '''
    async with db.execute(query_count, (language_from, letter.upper())) as cursor:
        total = (await cursor.fetchone())[0]
    
    # Получаем слова
    query = '''
        SELECT card_id, word, translation, card_type, language_from 
        FROM cards 
        WHERE language_from = ? AND first_letter = ?
        ORDER BY word ASC
        LIMIT ? OFFSET ?
    '''
    async with db.execute(query, (language_from, letter.upper(), per_page, offset)) as cursor:
        rows = await cursor.fetchall()
        cards = []
        for row in rows:
            cards.append({
                'card_id': row[0],
                'word': row[1],
                'translation': row[2],
                'card_type': row[3],
                'language_from': row[4]
            })
        return cards, total

async def get_all_letters(language_from: str) -> List[str]:
    """Получить все первые буквы для языка"""
    db = await get_db()
    async with db.execute('''
        SELECT DISTINCT first_letter FROM cards 
        WHERE language_from = ? AND first_letter IS NOT NULL AND first_letter != ""
        ORDER BY first_letter
    ''', (language_from,)) as cursor:
        rows = await cursor.fetchall()
        return [row[0] for row in rows]

async def search_words_by_pattern(pattern: str, page: int = 0, per_page: int = 5, language_from: str = None) -> Tuple[List[Dict], int]:
    """Поиск слов по шаблону с * (по слову, переводу и альтернативным переводам) с пагинацией"""
    db = await get_db()
    offset = page * per_page
    
    # Обработка шаблона с *
    if '*' in pattern:
        search_term = pattern.replace('*', '%')
    else:
        search_term = f'%{pattern}%'
    
    # Базовый запрос для подсчёта
    base_query_count = '''
        SELECT COUNT(*) FROM cards 
        WHERE word LIKE ? 
           OR translation LIKE ? 
           OR alternative_translations LIKE ?
    '''
    base_params = [search_term, search_term, f'%{search_term}%']
    
    # Базовый запрос для получения данных
    base_query = '''
        SELECT card_id, word, translation, card_type, language_from 
        FROM cards 
        WHERE word LIKE ? 
           OR translation LIKE ? 
           OR alternative_translations LIKE ?
        ORDER BY word ASC 
        LIMIT ? OFFSET ?
    '''
    base_params_data = [search_term, search_term, f'%{search_term}%', per_page, offset]
    
    # Если указан язык, добавляем условие
    if language_from:
        query_count = base_query_count + ' AND language_from = ?'
        params_count = base_params + [language_from]
        
        query = base_query.replace('ORDER BY', 'AND language_from = ? ORDER BY')
        params_data = base_params_data[:-2] + [language_from] + base_params_data[-2:]
    else:
        query_count = base_query_count
        params_count = base_params
        query = base_query
        params_data = base_params_data
    
    # Получаем общее количество
    async with db.execute(query_count, params_count) as cursor:
        total = (await cursor.fetchone())[0]
    
    # Получаем данные
    async with db.execute(query, params_data) as cursor:
        rows = await cursor.fetchall()
        cards = []
        for row in rows:
            cards.append({
                'card_id': row[0],
                'word': row[1],
                'translation': row[2],
                'card_type': row[3],
                'language_from': row[4]
            })
        return cards, total

# ============================================
# ФУНКЦИИ ДЛЯ ПОДСКАЗОК
# ============================================

async def get_hints_for_card(card_id: int, correct_answer: str, card_type: str, language_to: str):
    """Получить варианты для подсказки (для обычных слов)"""
    db = await get_db()
    async with db.execute('''
        SELECT translation FROM cards 
        WHERE card_type = ? AND language_to = ? AND card_id != ?
        ORDER BY RANDOM() 
        LIMIT 3
    ''', (card_type, language_to, card_id)) as cursor:
        wrong_hints = await cursor.fetchall()
        wrong_hints = [h[0] for h in wrong_hints if h[0] != correct_answer]
        
        # Если не хватает вариантов, добавляем заглушки
        while len(wrong_hints) < 3:
            wrong_hints.append(f"Вариант {len(wrong_hints) + 1}")
        
        # Перемешиваем все варианты
        all_hints = [correct_answer] + wrong_hints[:3]
        import random
        random.shuffle(all_hints)
        
        return [(hint, hint == correct_answer) for hint in all_hints]

async def get_multiple_choice_hints(card_id: int, correct_answer: str):
    """Получить варианты для идиом и типов без подсказки"""
    card = await get_card_by_id(card_id)
    if not card:
        return [(correct_answer, True), ("Вариант 1", False), ("Вариант 2", False), ("Вариант 3", False)]
    
    wrong_hints = card.get('wrong_hints', [])
    
    # Если неправильных вариантов меньше 3, добираем из базы
    while len(wrong_hints) < 3:
        db = await get_db()
        async with db.execute('''
            SELECT translation FROM cards 
            WHERE card_id != ? AND card_type = ? 
            ORDER BY RANDOM() 
            LIMIT 1
        ''', (card_id, card['card_type'])) as cursor:
            result = await cursor.fetchone()
            if result and result[0] != correct_answer and result[0] not in wrong_hints:
                wrong_hints.append(result[0])
            else:
                wrong_hints.append(f"Вариант {len(wrong_hints) + 1}")
    
    # Перемешиваем все варианты
    all_hints = [correct_answer] + wrong_hints[:3]
    import random
    random.shuffle(all_hints)
    
    return [(hint, hint == correct_answer) for hint in all_hints]

# ============================================
# ФУНКЦИИ ДЛЯ БЛАГОДАРНОСТИ
# ============================================

async def get_thanks_text() -> str:
    """Получить текст благодарности"""
    from config import DEFAULT_THANKS_TEXT
    db = await get_db()
    async with db.execute('SELECT value FROM bot_config WHERE key = "thanks_text"') as cursor:
        result = await cursor.fetchone()
        if result:
            return result[0]
        return DEFAULT_THANKS_TEXT

async def set_thanks_text(text: str):
    """Обновить текст благодарности (только для админа)"""
    db = await get_db()
    await db.execute('''
        UPDATE bot_config SET value = ?, updated_at = CURRENT_TIMESTAMP
        WHERE key = "thanks_text"
    ''', (text,))
    await db.commit()
    await auto_backup()

# ============================================
# ИМПОРТ ИЗ EXCEL — С АВТОБЭКАПОМ
# ============================================

def import_words_from_excel(file_path: str, user_id: int) -> tuple:
    """
    Импорт слов из Excel файла
    Возвращает: (количество добавленных, список ошибок)
    """
    import sqlite3
    import openpyxl
    import json
    import subprocess
    import os
    from config import CARD_TYPES
    
    errors = []
    added_count = 0
    
    try:
        workbook = openpyxl.load_workbook(file_path)
        sheet = workbook.active
        
        conn = sqlite3.connect("words.db")
        cursor = conn.cursor()
        
        for row in sheet.iter_rows(min_row=2, values_only=True):
            try:
                # Проверяем обязательные поля
                if not row[0] or not row[1]:
                    errors.append("Пропущена строка: пустое слово или перевод")
                    continue
                
                word = str(row[0]).strip()
                translation = str(row[1]).strip()
                language_from = str(row[2]).strip() if row[2] else 'mdf'
                language_to = str(row[3]).strip() if row[3] else 'ru'
                card_type = str(row[4]).strip() if row[4] else 'noun'
                image_url = str(row[5]).strip() if row[5] else None
                wrong_hints_text = str(row[6]).strip() if row[6] else ''
                topics_text = str(row[7]).strip() if row[7] else ''
                alt_trans_text = str(row[8]).strip() if row[8] else ''
                first_letter = word[0].upper() if word else None
                
                # Валидация
                if language_from not in ['ru', 'mdf']:
                    errors.append(f"Неверный язык: {language_from}. Пропускаем слово '{word}'")
                    continue
                if language_to not in ['ru', 'mdf']:
                    errors.append(f"Неверный язык: {language_to}. Пропускаем слово '{word}'")
                    continue
                if card_type not in CARD_TYPES:
                    errors.append(f"Неверный тип: {card_type}. Пропускаем слово '{word}'")
                    continue
                
                # Парсим неправильные варианты (в кавычках)
                wrong_hints = []
                if wrong_hints_text:
                    parts = []
                    current = []
                    in_quotes = False
                    for char in wrong_hints_text:
                        if char == '"':
                            in_quotes = not in_quotes
                            current.append(char)
                        elif char == ',' and not in_quotes:
                            parts.append(''.join(current).strip())
                            current = []
                        else:
                            current.append(char)
                    if current:
                        parts.append(''.join(current).strip())
                    wrong_hints = [p.strip('" ') for p in parts if p.strip()]
                
                # Парсим тематики и альтернативные переводы
                topics = [t.strip() for t in topics_text.split(',') if t.strip()] if topics_text else []
                alt_trans = [t.strip() for t in alt_trans_text.split(',') if t.strip()] if alt_trans_text else []
                
                # Проверяем, есть ли уже такое слово
                cursor.execute('SELECT card_id FROM cards WHERE word = ? AND language_from = ?',
                             (word, language_from))
                existing = cursor.fetchone()
                
                if existing:
                    # Обновляем существующее
                    wrong_hints_json = json.dumps(wrong_hints) if wrong_hints else None
                    topics_str = ','.join(topics) if topics else None
                    alt_trans_str = ','.join(alt_trans) if alt_trans else None
                    
                    cursor.execute('''
                        UPDATE cards 
                        SET translation = ?, language_to = ?, card_type = ?, 
                            image_url = ?, wrong_hints = ?, topics = ?,
                            alternative_translations = ?, first_letter = ?, 
                            updated_at = CURRENT_TIMESTAMP
                        WHERE card_id = ?
                    ''', (translation, language_to, card_type, image_url, 
                          wrong_hints_json, topics_str, alt_trans_str, 
                          first_letter, existing[0]))
                    added_count += 1
                    continue
                
                # Добавляем новое слово
                wrong_hints_json = json.dumps(wrong_hints) if wrong_hints else None
                topics_str = ','.join(topics) if topics else None
                alt_trans_str = ','.join(alt_trans) if alt_trans else None
                
                cursor.execute('''
                    INSERT INTO cards 
                    (word, translation, language_from, language_to, 
                     image_url, card_type, wrong_hints, topics, 
                     alternative_translations, first_letter, created_by)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ''', (word, translation, language_from, language_to, 
                      image_url, card_type, wrong_hints_json, topics_str, 
                      alt_trans_str, first_letter, user_id))
                added_count += 1
                
            except Exception as e:
                errors.append(f"Ошибка в строке: {str(e)}")
                continue
        
        conn.commit()
        conn.close()
        
        # Создаём бэкап после импорта
        if added_count > 0:
            try:
                subprocess.run(
                    ["python", "backup_db.py"],
                    capture_output=True,
                    text=True,
                    cwd=os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
                )
                print("✅ Автоматический бэкап после импорта создан")
            except:
                pass
        
    except Exception as e:
        errors.append(f"Ошибка при открытии файла: {str(e)}")
    
    return added_count, errors

# ============================================
# ФУНКЦИИ ДЛЯ ПРОСМОТРА ТАБЛИЦ
# ============================================

async def get_table_data(table_name: str, page: int = 0, per_page: int = 20) -> dict:
    """Получить данные из таблицы SQLite с пагинацией"""
    import sqlite3
    try:
        conn = sqlite3.connect("words.db")
        cursor = conn.cursor()
        
        # Проверяем, существует ли таблица
        cursor.execute("SELECT name FROM sqlite_master WHERE type='table' AND name=?", (table_name,))
        if not cursor.fetchone():
            conn.close()
            return {"error": f"Таблица '{table_name}' не найдена"}
        
        # Получаем структуру таблицы
        cursor.execute(f"PRAGMA table_info({table_name})")
        columns = [col[1] for col in cursor.fetchall()]
        
        # Получаем общее количество записей
        cursor.execute(f"SELECT COUNT(*) FROM {table_name}")
        total = cursor.fetchone()[0]
        
        # Получаем данные с пагинацией
        offset = page * per_page
        cursor.execute(f"SELECT * FROM {table_name} LIMIT ? OFFSET ?", (per_page, offset))
        rows = cursor.fetchall()
        
        conn.close()
        
        return {
            "columns": columns,
            "rows": rows,
            "total": total,
            "per_page": per_page,
            "page": page,
            "table_name": table_name
        }
    except Exception as e:
        return {"error": str(e)}

async def format_table_message(data: dict) -> str:
    """Форматировать данные таблицы в читаемое сообщение"""
    if "error" in data:
        return f"❌ {data['error']}"
    
    table_name = data['table_name']
    columns = data['columns']
    rows = data['rows']
    total = data['total']
    per_page = data['per_page']
    page = data['page']
    
    if not rows and total == 0:
        return f"📭 Таблица '{table_name}' пуста"
    
    total_pages = (total + per_page - 1) // per_page if total > 0 else 1
    start_record = page * per_page + 1
    end_record = min((page + 1) * per_page, total)
    
    message = f"📊 <b>Таблица: {table_name}</b>\n"
    message += f"📝 Всего записей: {total}\n"
    message += f"📄 Показано: {start_record}-{end_record} (страница {page+1}/{total_pages})\n\n"
    
    # Заголовки
    message += "│ " + " │ ".join(columns) + " │\n"
    message += "├" + "─" * (len(" │ ".join(columns)) + 2) + "┤\n"
    
    # Данные
    for row in rows:
        row_list = list(row)
        for j, col in enumerate(columns):
            # Если это JSON-поле, пытаемся распарсить
            if col in ['wrong_hints'] and row_list[j]:
                try:
                    row_list[j] = json.loads(row_list[j])
                except:
                    pass
            # Обрезаем длинные строки
            if isinstance(row_list[j], str) and len(row_list[j]) > 30:
                row_list[j] = row_list[j][:27] + "..."
        
        message += "│ " + " │ ".join(str(val) for val in row_list) + " │\n"
    
    return message